"""Backtest result models and file exports."""

import csv
import json
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path

from backtests.metrics import PerformanceMetrics
from models.closed_trade import ClosedTrade
from models.terminated_setup import TerminatedSetup


@dataclass(frozen=True)
class SignalOutcome:
    """A generated signal and an optional realized research outcome."""

    symbol: str
    timestamp: datetime
    signal_type: str
    zone_id: int
    price: float
    outcome: float | None = None

    def to_dict(self) -> dict[str, object]:
        payload = asdict(self)
        payload["timestamp"] = self.timestamp.isoformat()
        return payload


@dataclass(frozen=True)
class BacktestResult:
    """Canonical measurable output from one strategy/dataset run."""

    strategy_id: str
    strategy_name: str
    symbols: tuple[str, ...]
    started_at: datetime | None
    ended_at: datetime | None
    total_signals: int
    wins: int
    losses: int
    win_rate: float
    profit_factor: float | None
    average_win: float
    average_loss: float
    expectancy: float
    max_drawdown: float
    largest_win: float
    largest_loss: float
    signal_outcomes: tuple[SignalOutcome, ...] = field(default_factory=tuple)
    closed_trades: tuple[ClosedTrade, ...] = field(default_factory=tuple)
    terminated_setups: tuple[TerminatedSetup, ...] = field(default_factory=tuple)
    total_trades: int = 0
    consecutive_wins: int = 0
    consecutive_losses: int = 0
    ambiguous_exits: int = 0
    ambiguous_rate: float = 0.0
    avg_holding_duration: float = 0.0
    expired_setups_count: int = 0
    cancelled_setups_count: int = 0
    rejected_setups_count: int = 0
    ignored_active_lifecycle_count: int = 0
    rejected_policy_boundary_count: int = 0
    accepted_setups_count: int = 0
    currently_active_count: int = 0

    @classmethod
    def from_metrics(
        cls,
        *,
        strategy_id: str,
        strategy_name: str,
        symbols: tuple[str, ...],
        started_at: datetime | None,
        ended_at: datetime | None,
        metrics: PerformanceMetrics,
        signal_outcomes: tuple[SignalOutcome, ...],
        closed_trades: tuple[ClosedTrade, ...] = (),
        terminated_setups: tuple[TerminatedSetup, ...] = (),
        ignored_active_lifecycle_count: int = 0,
        rejected_policy_boundary_count: int = 0,
        accepted_setups_count: int = 0,
        currently_active_count: int = 0,
    ) -> "BacktestResult":
        expired = sum(1 for s in terminated_setups if s.reason == "EXPIRED")
        cancelled = sum(1 for s in terminated_setups if s.reason == "CANCELLED_END_OF_DATA")
        rejected = sum(1 for s in terminated_setups if s.reason == "REJECTED_INVALID_ENTRY")
        return cls(
            strategy_id=strategy_id,
            strategy_name=strategy_name,
            symbols=symbols,
            started_at=started_at,
            ended_at=ended_at,
            signal_outcomes=signal_outcomes,
            closed_trades=closed_trades,
            terminated_setups=terminated_setups,
            expired_setups_count=expired,
            cancelled_setups_count=cancelled,
            rejected_setups_count=rejected,
            ignored_active_lifecycle_count=ignored_active_lifecycle_count,
            rejected_policy_boundary_count=rejected_policy_boundary_count,
            accepted_setups_count=accepted_setups_count,
            currently_active_count=currently_active_count,
            **asdict(metrics),
        )

    def reconcile_signals(self) -> dict[str, int]:
        generated = self.total_signals
        accepted = self.accepted_setups_count
        ignored = self.ignored_active_lifecycle_count
        rejected = self.rejected_policy_boundary_count
        diff = generated - (accepted + ignored + rejected)
        return {
            "generated": generated,
            "accepted": accepted,
            "ignored": ignored,
            "rejected": rejected,
            "difference": diff,
        }

    def reconcile_execution(self) -> dict[str, int]:
        accepted = self.accepted_setups_count
        closed = self.total_trades
        expired = self.expired_setups_count
        rejected = self.rejected_setups_count
        cancelled = self.cancelled_setups_count
        active = self.currently_active_count
        diff = accepted - (closed + expired + rejected + cancelled + active)
        return {
            "accepted": accepted,
            "closed": closed,
            "expired": expired,
            "rejected": rejected,
            "cancelled": cancelled,
            "active": active,
            "difference": diff,
        }

    def summary_dict(self) -> dict[str, object]:
        return {
            "strategy_id": self.strategy_id,
            "strategy_name": self.strategy_name,
            "symbols": ",".join(self.symbols),
            "started_at": self.started_at.isoformat() if self.started_at else None,
            "ended_at": self.ended_at.isoformat() if self.ended_at else None,
            "total_signals": self.total_signals,
            "wins": self.wins,
            "losses": self.losses,
            "win_rate": self.win_rate,
            "profit_factor": self.profit_factor,
            "average_win": self.average_win,
            "average_loss": self.average_loss,
            "expectancy": self.expectancy,
            "max_drawdown": self.max_drawdown,
            "largest_win": self.largest_win,
            "largest_loss": self.largest_loss,
            "total_trades": self.total_trades,
            "consecutive_wins": self.consecutive_wins,
            "consecutive_losses": self.consecutive_losses,
            "ambiguous_exits": self.ambiguous_exits,
            "ambiguous_rate": self.ambiguous_rate,
            "avg_holding_duration": self.avg_holding_duration,
            "expired_setups_count": self.expired_setups_count,
            "cancelled_setups_count": self.cancelled_setups_count,
            "rejected_setups_count": self.rejected_setups_count,
            "ignored_active_lifecycle_count": self.ignored_active_lifecycle_count,
            "rejected_policy_boundary_count": self.rejected_policy_boundary_count,
            "accepted_setups_count": self.accepted_setups_count,
            "currently_active_count": self.currently_active_count,
        }

    def to_dict(self) -> dict[str, object]:
        payload = self.summary_dict()
        payload["symbols"] = list(self.symbols)
        payload["signal_outcomes"] = [
            outcome.to_dict() for outcome in self.signal_outcomes
        ]
        payload["closed_trades"] = [
            asdict(trade) for trade in self.closed_trades
        ]
        for trade in payload["closed_trades"]:
            if isinstance(trade["entry_time"], datetime):
                trade["entry_time"] = trade["entry_time"].isoformat()
            if isinstance(trade["exit_time"], datetime):
                trade["exit_time"] = trade["exit_time"].isoformat()
                
        payload["terminated_setups"] = [
            asdict(setup) for setup in self.terminated_setups
        ]
        for s in payload["terminated_setups"]:
            if isinstance(s["timestamp"], datetime):
                s["timestamp"] = s["timestamp"].isoformat()
            if isinstance(s["termination_time"], datetime):
                s["termination_time"] = s["termination_time"].isoformat()
        return payload

    def export_json(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(self.to_dict(), indent=2), encoding="utf-8")
        return path

    def export_csv(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", newline="", encoding="utf-8") as file:
            writer = csv.DictWriter(file, fieldnames=list(self.summary_dict()))
            writer.writeheader()
            writer.writerow(self.summary_dict())
        return path

    def export_excel(self, path: Path) -> Path:
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise RuntimeError(
                "Excel export requires openpyxl. Install project requirements."
            ) from exc

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        for row_index, (key, value) in enumerate(
            self.summary_dict().items(), start=1
        ):
            summary.cell(row=row_index, column=1, value=key)
            summary.cell(row=row_index, column=2, value=value)

        outcomes = workbook.create_sheet("Signal Outcomes")
        headers = ["symbol", "timestamp", "signal_type", "zone_id", "price", "outcome"]
        outcomes.append(headers)
        for record in self.signal_outcomes:
            payload = record.to_dict()
            outcomes.append([payload[header] for header in headers])

        if self.closed_trades:
            trades_sheet = workbook.create_sheet("Closed Trades")
            trade_headers = ["symbol", "direction", "entry_price", "entry_time", "exit_price", "exit_time", "sl_price", "tp_price", "pnl_r", "close_reason", "holding_duration_candles"]
            trades_sheet.append(trade_headers)
            for trade in self.closed_trades:
                trades_sheet.append([
                    trade.symbol,
                    trade.direction,
                    trade.entry_price,
                    trade.entry_time.isoformat() if isinstance(trade.entry_time, datetime) else str(trade.entry_time),
                    trade.exit_price,
                    trade.exit_time.isoformat() if isinstance(trade.exit_time, datetime) else str(trade.exit_time),
                    trade.sl_price,
                    trade.tp_price,
                    trade.pnl_r,
                    trade.close_reason,
                    trade.holding_duration_candles
                ])

        workbook.save(path)
        return path

